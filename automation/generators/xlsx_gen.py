import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models.schemas import XlsxRequest

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "generated_files")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def generate_xlsx(req: XlsxRequest) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = req.sheet_name or "Dados"

    # Color definitions
    navy_fill = PatternFill(start_color="0A1128", end_color="0A1128", fill_type="solid")
    cyan_accent = PatternFill(start_color="00B4D8", end_color="00B4D8", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0A1128")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    data_font = Font(name="Calibri", size=10, color="1E293B")
    total_font = Font(name="Calibri", size=11, bold=True, color="0A1128")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    cyan_bottom = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='medium', color='00B4D8')
    )

    # Row 1: Title
    ws.cell(row=1, column=1, value=req.title or "Verity AI — Relatório de Dados").font = title_font
    ws.row_dimensions[1].height = 28

    # Row 2: Subtitle
    ws.cell(row=2, column=1, value="Gerado automaticamente pela Verity AI • www.verityai.com").font = subtitle_font
    ws.row_dimensions[2].height = 18

    # Row 4: Headers
    start_row = 4
    num_cols = len(req.headers)
    for col_idx, header in enumerate(req.headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = white_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cyan_bottom
    ws.row_dimensions[start_row].height = 26

    # Data Rows
    current_row = start_row + 1
    numeric_columns = set()

    for r_idx, row_data in enumerate(req.rows):
        is_zebra = (r_idx % 2 == 1)
        ws.row_dimensions[current_row].height = 20

        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=c_idx)
            
            # Type casting if string looks like number
            parsed_val = val
            if isinstance(val, str):
                cleaned = val.replace("R$", "").replace("$", "").replace(".", "").replace(",", ".").strip()
                try:
                    if "." in cleaned:
                        parsed_val = float(cleaned)
                    else:
                        parsed_val = int(cleaned)
                except ValueError:
                    parsed_val = val

            cell.value = parsed_val
            cell.font = data_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

            if isinstance(parsed_val, (int, float)):
                numeric_columns.add(c_idx)
                cell.alignment = Alignment(horizontal="right", vertical="center")
                # Format currency / decimals
                if isinstance(parsed_val, float):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Summary row
    if req.include_summary_row and req.rows and numeric_columns:
        ws.row_dimensions[current_row].height = 22
        first_cell = ws.cell(row=current_row, column=1, value="TOTAL")
        first_cell.font = total_font
        first_cell.border = thin_border

        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if col_idx in numeric_columns:
                col_letter = get_column_letter(col_idx)
                cell.value = f"=SUM({col_letter}5:{col_letter}{current_row - 1})"
                cell.font = total_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            elif col_idx != 1:
                cell.value = ""

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save
    clean_name = os.path.basename(req.filename or "planilha_verity.xlsx")
    if not clean_name.endswith(".xlsx"):
        clean_name += ".xlsx"
    filepath = os.path.join(OUTPUT_DIR, clean_name)

    wb.save(filepath)
    file_size = os.path.getsize(filepath)

    return {
        "success": True,
        "filename": clean_name,
        "file_size": file_size,
        "download_url": f"/api/v1/download/{clean_name}",
        "message": f"Planilha Excel '{clean_name}' gerada com sucesso com formatação corporativa e fórmulas."
    }
